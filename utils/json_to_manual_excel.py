import argparse
import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation


UNKNOWN_MARKERS = {
    "unknown",
    "not known",
    "not answerable",
    "cannot answer",
    "cannot be determined",
    "cant answer",
    "cant be determined",
    "can't answer",
    "can't be determined",
    "undetermined",
    "not enough info",
    "not enough information",
}


def load_json(path: Path):
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def join_list(values):
    if isinstance(values, list):
        return "; ".join(str(v) for v in values)
    return "" if values is None else str(values)


def to_json_string(value):
    return json.dumps(value, ensure_ascii=False)


def get_answer_text(record, ans_key):
    answers = record.get("answers", {})
    if isinstance(answers, dict):
        return answers.get(ans_key, "")
    return ""


def get_answer_info_field(record, ans_key, index):
    answer_info = record.get("answer_info", {})
    value = answer_info.get(ans_key, ["", ""])

    if isinstance(value, list) and len(value) > index:
        return value[index]

    return ""


def normalize_unknown_text(text):
    if text is None:
        return ""

    text = str(text).lower().strip()
    text = text.replace("'", "")
    text = " ".join(text.split())
    return text


def is_unknown_value(value):
    value_norm = normalize_unknown_text(value)
    markers_norm = {normalize_unknown_text(m) for m in UNKNOWN_MARKERS}
    return value_norm in markers_norm


def detect_unknown_letter(record):
    answer_info = record.get("answer_info", {})

    mapping = {
        "ans0": "A",
        "ans1": "B",
        "ans2": "C",
    }

    candidates = []

    for ans_key, letter in mapping.items():
        value = answer_info.get(ans_key, ["", ""])

        field0 = value[0] if isinstance(value, list) and len(value) > 0 else ""
        field1 = value[1] if isinstance(value, list) and len(value) > 1 else ""

        if is_unknown_value(field0) or is_unknown_value(field1):
            candidates.append(letter)

    if len(candidates) == 1:
        return candidates[0]

    return ""


def build_rows(records):
    rows = []

    for record in records:
        row = {
            "example_id": record.get("example_id", ""),
            "category": record.get("category", ""),
            "question_polarity": record.get("question_polarity", ""),
            "stereotyped_groups": join_list(record.get("stereotyped_groups", [])),

            "context": record.get("context", ""),
            "question": record.get("question", ""),
            "label": record.get("label", ""),
            "gold_answer": record.get("gold_answer", ""),

            "A_answer": get_answer_text(record, "ans0"),
            "A_info_0": get_answer_info_field(record, "ans0", 0),
            "A_info_1": get_answer_info_field(record, "ans0", 1),

            "B_answer": get_answer_text(record, "ans1"),
            "B_info_0": get_answer_info_field(record, "ans1", 0),
            "B_info_1": get_answer_info_field(record, "ans1", 1),

            "C_answer": get_answer_text(record, "ans2"),
            "C_info_0": get_answer_info_field(record, "ans2", 0),
            "C_info_1": get_answer_info_field(record, "ans2", 1),

            # Automatically pre-filled from answer_info.
            "unknown": detect_unknown_letter(record),

            # Columns to be manually annotated.
            "target": "",
            "non_target": "",
        }

        rows.append(row)

    return rows


def export_excel(input_json: Path, output_xlsx: Path):
    records = load_json(input_json)
    rows = build_rows(records)

    df = pd.DataFrame(rows)

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="annotation")

    format_excel(output_xlsx)

    print(f"[OK] Excel created: {output_xlsx}")
    print(f"Exported records: {len(df)}")
    print(f"Pre-filled unknown values: {df['unknown'].astype(bool).sum()}/{len(df)}")


def format_excel(path: Path):
    wb = load_workbook(path)
    ws = wb["annotation"]

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    manual_fill = PatternFill("solid", fgColor="FFF2CC")
    auto_fill = PatternFill("solid", fgColor="E2F0D9")

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    header_to_col = {cell.value: cell.column for cell in ws[1]}

    # The unknown column is automatically filled.
    for col_name in ["unknown"]:
        col_idx = header_to_col.get(col_name)
        if col_idx is not None:
            for row_idx in range(1, ws.max_row + 1):
                ws.cell(row=row_idx, column=col_idx).fill = auto_fill

    # The target and non_target columns must be manually filled.
    for col_name in ["target", "non_target"]:
        col_idx = header_to_col.get(col_name)
        if col_idx is not None:
            for row_idx in range(1, ws.max_row + 1):
                ws.cell(row=row_idx, column=col_idx).fill = manual_fill

    letter_validation = DataValidation(
        type="list",
        formula1='"A,B,C"',
        allow_blank=True,
    )
    ws.add_data_validation(letter_validation)

    for col_name in ["unknown", "target", "non_target"]:
        col_idx = header_to_col.get(col_name)
        if col_idx is not None:
            col_letter = ws.cell(row=1, column=col_idx).column_letter
            letter_validation.add(f"{col_letter}2:{col_letter}{ws.max_row}")

    column_widths = {
        "A": 12,
        "B": 24,
        "C": 18,
        "D": 35,
        "E": 70,
        "F": 45,
        "G": 10,
        "H": 35,

        "I": 35,
        "J": 28,
        "K": 24,

        "L": 35,
        "M": 28,
        "N": 24,

        "O": 35,
        "P": 28,
        "Q": 24,

        "R": 12,
        "S": 12,
        "T": 12,

        "U": 90,
    }

    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 70

    wb.save(path)


def main():
    parser = argparse.ArgumentParser(
        description="Convert BBQ JSON to a flat Excel file for manual target/non-target annotation."
    )
    parser.add_argument(
        "--input-json",
        type=Path,
        default=Path("data/processed/bbq_disambiguated_sample_380.json"),
        help="Input JSON file.",
    )
    parser.add_argument(
        "--output-xlsx",
        type=Path,
        default=Path("data/processed/bbq_disambiguated_sample_380_manual_annotation.xlsx"),
        help="Output Excel file.",
    )

    args = parser.parse_args()

    export_excel(args.input_json, args.output_xlsx)


if __name__ == "__main__":
    main()